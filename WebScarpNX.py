import urllib2
from bs4 import BeautifulSoup
from docx import Document

def remove_row(table, row):
    tbl = table._tbl
    tr = row._tr
    tbl.remove(tr)


DocX        = Document()
url         = "file:///C:/Users/Sunshine/Desktop/90.html" 
page        = urllib2.urlopen(url).read()
soup        = BeautifulSoup(page,"lxml")
TopRow      = ["Ser","Alert Type","ID","Malware","Time","Source IP","URL","Location","Badges"]
SelectedCol = [1,2,3,5,7,8,10,13,14]
TableRow    = (len(soup.find_all('tr')))-1
TableCol    = 9
ListDelete  = []
DocTable    = DocX.add_table(rows=TableRow,cols=TableCol,style='Table Grid')
ColumnsHandle              = 0
RowsToFill                 = 1
headlen                    = 0
#=================HEADINGS====================
#Cell=DocTable.cell(0,0)

for var in TopRow:
    Cell=DocTable.cell(0,headlen)
    Cell.text = var
    headlen += 1
DocX.save('C:\Users\Sunshine\Desktop\RESULTS.docx')
#==============VALUES============================
for tr in soup.find_all('tr')[1:]:
    if ( RowsToFill  < TableRow ):  
                ColumnsHandle = 1
                tds = tr.find_all('td')
                for col in range (1 , TableCol):
                        DocX.save('C:\Users\Sunshine\Desktop\RESULTS.docx')
                        if (ColumnsHandle < len(SelectedCol)):
                            Cell=DocTable.cell(RowsToFill ,col)
                            Cell.text = tds[SelectedCol[ColumnsHandle]].text
                            if (ColumnsHandle == 5 ):
                                if ( (DocTable.cell(RowsToFill,col).text) == (DocTable.cell(RowsToFill-1,col).text) ):
                                    ListDelete.append(RowsToFill)
                            ColumnsHandle += 1
                RowsToFill  +=1 
for rowvar in range (len(ListDelete)-2,-1 ,-1):
    row = DocTable.rows[ListDelete[rowvar]]
    remove_row(DocTable, row)

for num in range (1,TableRow-len(ListDelete)+1):
    Cell=DocTable.cell(num,0)
    Cell.text = str(num)
DocX.save('C:\Users\Sunshine\Desktop\RESULTS.docx')


#===============NOT TO BE Deleted=======================

'''
print "Alert Type : %s | ID: %s | Malware: %s | Time: %s | Source IP: %s | URL: %s | Location: %s | Badges: %s " % \
                        (tds[2].text.strip(),tds[3].text,tds[5].text,tds[7].text,tds[8].text,tds[10].text,tds[13].text,tds[14].text.strip())

'''
